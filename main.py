'''
# OTDScript_V1 -Script для ОТДЕЛА ТЕХНИЧЕСКОЙ ДОКУМЕНТАЦИИ версия 1
# СКРИПТ ДЛЯ АВТОЗАПОЛНЕНИЯ ДОКУМЕНТОВ ДЛЯ ОТДЕЛА ТЕХНИЧЕСКОЙ ДОКУМЕНТАЦИИ

Цели:
    -Автоматизировать однотипные задачи отдела.
    -Сократить время обработки документов

Бизнес Требования к скрипту:
    - Скрипт должен читать INPUT данные
    - Записывать их на нужные места в документах
    - Сохраниять документ с записанными данными

    
План выполнения работ:

1. Нстроить шаблоны
    - шблон "1 Журнал авторского надзора"
    - шблон "2 Журнал сварочных работ"
    - шблон "3 Акт скрытых работ"
    - шблон "4 Журнал работ по МСК"
    - шблон "5 Журнал производства работ"
2. Написать скрипт который прочитает данные из Excel
3. Написать скрипт заполняющий файлы данными:
    - Скрипт для "1 Журнал авторского надзора"
    - Скрипт для "2 Журнал сварочных работ"
    - Скрипт для "3 Акт скрытых работ"
    - Скрипт для "4 Журнал работ по МСК"
    - Скрипт для "5 Журнал производства работ"
4. Тестирование
5. Стилизация
6. Конвертирование в exe файл
7. Тест в формате EXE
8. Обучение пользователей
9. Доработка
10. Отчет

'''

import openpyxl
from datetime import datetime
from docxtpl import DocxTemplate
from pyautogui import alert
import os

TEMPLATES_FOLDER_PATH = "Шаблоны/"
OFP = ""
INPUT_FILE_PATH = "INPUT.xlsx"
INPUT_FILE_PAGE_NAME = "Основные переменные"
obj_nameanius = "БС"

TEMPLATES = [
    { 
        "TEMPLATE" : TEMPLATES_FOLDER_PATH+"1 Журнал авторского надзора.docx",
        "FILE_NAME" :     "1Журнал авторского надзора"
    },{ 
        "TEMPLATE" : TEMPLATES_FOLDER_PATH+"2 Журнал сварочных работ.docx",
        "FILE_NAME" :     "2Журнал сварочных работ"
    },
    { 
        "TEMPLATE" : TEMPLATES_FOLDER_PATH+"3 Акт скрытых работ.docx",
        "FILE_NAME" :     "3Акт скрытых работ"
    },
    { 
        "TEMPLATE" : TEMPLATES_FOLDER_PATH+"4 Журнал работ по МСК.docx",
        "FILE_NAME" : "4Журнал работ по МСК"
    },{ 
        "TEMPLATE" : TEMPLATES_FOLDER_PATH+"5 Журнал производства работ.docx",
        "FILE_NAME" : "5Журнал производства работ"
    }
]

OUTPUT_FILE_NAMES = [
    "Журнал авторского надзора",
    "Журнал сварочных работ",
    "Акт скрытых работ",
    "Журнал работ по МСК",
    "Журнал производства работ"
]


def none_to_empty_string(dictionary):
    return {key: value if value is not None else "" for key, value in dictionary.items()}


def get_excel_data():
    try:
    # if True:
        # Открыть файл и взять данные из нужного листа 
        sheet = openpyxl.load_workbook(INPUT_FILE_PATH, data_only=True)[INPUT_FILE_PAGE_NAME]
        
       # Присваивание в data_dict значение из листа 
        data_dict = {}
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if len(row) >= 2:
                key, value = row[:2]
                data_dict[key] = value


        # Удаление лишних строк
        for i in range(1,16):
            if data_dict[f'subcontractor_{i}'] == None:
                data_dict[f'DEFAULT_VALUE_{i}'] = None
                data_dict[f'CSJ_ORG_{i}'] = None
                data_dict[f'ne_bilo_{i}'] = None
                data_dict[f'object_responsible_{i}'] = None
        for i in range(1,31):
            if data_dict[f'ASR__akt_whalders_{i}'] == None:
                data_dict[f'aktosrnumb_{i}'] = None

        # Превратить все None в пустой string ""
        data_dict = none_to_empty_string(data_dict)
        
        # Удалить название словаря None
        del data_dict[None]

        return data_dict
    # Обработчик ошибок
    except Exception as e:
        alert(f"Ошибка в Ексель файле\nError message: {e} (EXCEL)")
        return None


def replace_words_in_docx(dictionary, file_path, file_name, OUTPUT_FOLDER_PATH=OFP):
    try:
    # if True:
        current_date = datetime.now().strftime('%Y-%m-%d_%H-%M-%S')
        output_file_path = f"{OUTPUT_FOLDER_PATH}{file_name} {dictionary['object_name']} {current_date}.docx"

        # Берет шаблон, заполняет данными из словаря, сохраняет как новый файл
        print(file_path)
        doc = DocxTemplate(file_path)
        doc.render(dictionary)
        doc.save(output_file_path)

    except Exception as e:
        print(f"Ошибка в шаблоне Word файла\nError message: {e} (WORD)")
        alert(f"Ошибка в шаблоне Word файла\nError message: {e} (WORD)")
        return None





def preparatory_operations():
    pass


import os
from docx import Document



def merge_docx_files(input_files, output_file):
    merged_doc = Document()

    for file in input_files:
        doc = Document(file)

        for element in doc.element.body:
            merged_doc.element.body.append(element)

    merged_doc.save(output_file)
    print(f"Файлы успешно объединены и сохранены в '{output_file}'.")




def data_assignments():
    global obj_nameanius
    
    data = get_excel_data()
    current_date = datetime.now().strftime('_%d-%m-%Y')
    obj_nameanius = data['object_name']
    OUTPUT_FOLDER_PATH = data['object_name']+ f"{current_date}/"
        
    try:
        os.mkdir(OUTPUT_FOLDER_PATH)
    except:
        pass
    for tmpl in TEMPLATES:
        # alert(tmpl['FILE_NAME'])
        if tmpl['FILE_NAME'] == "3Акт скрытых работ":
            count = 0 
            for i in range(1,30):
                param_1 = data['ASR__akt_name_'+str(i)]
                param_2 = data['ASR__akt_num_'+str(i)]
                param_3 = data['ASR__akt_deskrop_'+str(i)]
                param_4 = data['ASR__akt_useditems_'+str(i)]

                param_5 = data['ASR__akt_startdate_'+str(i)]
                param_6 = data['ASR__akt_enddate_'+str(i)]
                param_7 = data['ASR__akt_name_'+str(i+1)]

                if param_7 == "":
                    param_7 = "Установка ФБС блоков и ЦМ контейнера"
                    count = i+1


                XXX = {"param_1":param_1,"param_2":param_2,"param_3":param_3,"param_4":param_4,"param_5":param_5,"param_6":param_6,"param_7":param_7}                   
                print(XXX)
                data.update(XXX)

                doc = DocxTemplate(TEMPLATES_FOLDER_PATH+"Акт скрытых работ/3 Акт скрытых работ.docx")
                doc.render(data)
                doc.save(TEMPLATES_FOLDER_PATH+f"Акт скрытых работ/file{i}.docx")
                if param_7 == "Установка ФБС блоков и ЦМ контейнера":
                    break
            
            file_names = []
            for i in range(1,count):
                file_names.append(TEMPLATES_FOLDER_PATH+f"Акт скрытых работ/file{i}.docx")
            merge_docx_files(file_names, f"{OUTPUT_FOLDER_PATH}{tmpl['FILE_NAME']} {data['object_name']} {current_date}.docx")
            
            file_names = []
            for i in range(1,30):
                file_names.append(TEMPLATES_FOLDER_PATH+f"Акт скрытых работ/file{i}.docx")
            for file_name in file_names:
                try:
                    os.remove(file_name)
                    print(f"Файл '{file_name}' успешно удален.")
                except FileNotFoundError:
                    print(f"Файл '{file_name}' не найден.")
                except Exception as e:
                    print(f"Произошла ошибка при удалении файла '{file_name}': {e}")
                
        else:
            replace_words_in_docx(data, tmpl['TEMPLATE'], tmpl['FILE_NAME'], OUTPUT_FOLDER_PATH)

if __name__ == "__main__":
    preparatory_operations()
    data_assignments()
    import reporter
    x = reporter.send_report( 
        process="Заполнение журнала", responsible="отдел Меруерт", text=obj_nameanius
    )
    print(x.text)
    alert("Конец процесса")





